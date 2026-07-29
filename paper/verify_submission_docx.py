"""Structural QA for the generated Crime and Policy submission-review DOCX."""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "paper" / "범죄와정책_투고원고_검증본.docx"
ANONYMOUS_DOCX = ROOT / "paper" / "범죄와정책_심사용_익명원고_최종.docx"


def dxa(value):
    return int(value / 635)


def verify_document(path, anonymous):
    document = Document(path)
    section = document.sections[0]
    text = "\n".join(p.text for p in document.paragraphs)
    normalized_text = text.replace("\u00a0", " ")

    expected = {
        "page_width": 11792,   # 208 mm
        "page_height": 15704,  # 277 mm
        "top_margin": 1134,    # 20 mm
        "bottom_margin": 850,  # 15 mm
        "left_margin": 850,
        "right_margin": 850,
        "header_distance": 850,
        "footer_distance": 850,
    }
    actual = {
        "page_width": dxa(section.page_width),
        "page_height": dxa(section.page_height),
        "top_margin": dxa(section.top_margin),
        "bottom_margin": dxa(section.bottom_margin),
        "left_margin": dxa(section.left_margin),
        "right_margin": dxa(section.right_margin),
        "header_distance": dxa(section.header_distance),
        "footer_distance": dxa(section.footer_distance),
    }
    for key, target in expected.items():
        assert abs(actual[key] - target) <= 2, (key, actual[key], target)

    assert len(document.sections) == 1
    assert len(document.tables) == 10
    assert document.core_properties.author == ""
    assert document.core_properties.last_modified_by == ""
    assert document.core_properties.title == (
        "생성형 AI의 한국어 개인식별정보(PII) 유출 위험과 "
        "정규화 기반 Layer 0의 필요성"
    )
    assert document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert document.paragraphs[0].runs[0].font.name == "HY견고딕"
    assert document.paragraphs[0].runs[0].font.size.pt == 15.0
    if anonymous:
        assert all("민우" not in paragraph.text for paragraph in document.paragraphs)
        assert all("CCIT" not in paragraph.text for paragraph in document.paragraphs)
        assert all("KPIIGD" not in paragraph.text for paragraph in document.paragraphs)
        assert all(
            "694ca717dd47e3d8f229bfa4da84c1fad607576b" not in paragraph.text
            for paragraph in document.paragraphs
        )
        assert document.core_properties.subject == "「범죄와 정책」 심사용 익명원고"
    else:
        assert document.paragraphs[3].text == "민우"
        assert document.paragraphs[3].alignment == WD_ALIGN_PARAGRAPH.RIGHT

    paragraph_texts = [paragraph.text for paragraph in document.paragraphs]
    required_order = ("〈요 약〉", "목 차", "Ⅰ. 서론", "참 고 문 헌", "ABSTRACT")
    positions = [paragraph_texts.index(marker) for marker in required_order]
    assert positions == sorted(positions), (required_order, positions)
    korean_abstract = paragraph_texts[paragraph_texts.index("〈요 약〉") + 1]
    english_abstract = paragraph_texts[paragraph_texts.index("ABSTRACT") + 2]
    assert 550 <= len(korean_abstract) <= 800, len(korean_abstract)
    assert 1000 <= len(english_abstract) <= 1600, len(english_abstract)
    assert document.paragraphs[paragraph_texts.index("Ⅰ. 서론") + 2].runs[
        0
    ].font.name == "HY신명조"
    assert document.paragraphs[
        paragraph_texts.index("Ⅰ. 서론")
    ].paragraph_format.page_break_before
    assert document.paragraphs[
        paragraph_texts.index("참 고 문 헌")
    ].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert document.paragraphs[
        paragraph_texts.index("〈요 약〉") + 1
    ].paragraph_format.first_line_indent
    assert document.paragraphs[
        paragraph_texts.index("ABSTRACT") + 2
    ].paragraph_format.first_line_indent

    required_text = (
        "생성형 AI의 한국어 개인식별정보(PII) 유출 위험과 정규화 기반 Layer 0의 필요성",
        "Korean Personally Identifiable Information Leakage Risks in Generative AI",
        "9,964건(99.64%)",
        "1,354건",
        "9.12%에서 18.75%",
        "Layer 0",
        "성별·젠더 변수를 포함하지 않은 시스템 감사",
        "A Lifestyle-Routine Activity Theory (LRAT) Approach",
        "다중 전자문서 환경에서 문서 구조 기반 개인정보 노출 위험과 비식별화 처리에 관한 연구",
        "백서진, 최한림, 박윤지, 정보남, 함근희, 2026",
        "OpenAI Codex",
        "독립적인 학술 근거나 연구자료로 채택하지 않았으며",
        "presidio.dataprivacystack.org/supported_entities/",
        "1,519개 중 29개(1.91%)",
        "693개 중 48개(6.93%)",
        "의료 265개 중 44개(16.60%)",
        "개발문서 255개 중 4개(1.57%)",
        "조문별 시행일은 국가법령정보센터 공식 XML에서 모두 2026년 1월 22일",
        "<표 1> 저장소별 감사 기준과 역할",
        "<표 9> Layer 0와 LLM 판별기의 지연시간",
    )
    for claim in required_text:
        assert claim in normalized_text, claim

    forbidden_text = (
        "변이 사례의 미차단율은 원형보다 3.93배",
        "2026년 7월 21일 시행된 「인공지능 발전과 신뢰 기반 조성 등에 관한 기본법」",
        "생성형 AI의 한국어 민감정보 유출 위험과 정규화 기반 Layer 0의 필요성",
        "Korean Sensitive-Information Leakage Risks in Generative AI",
        "이윤호·김도우·유영재",
        "박기태·조제성",
        "서지혜·최진선·박진수",
        "**",
        "```",
    )
    for claim in forbidden_text:
        assert claim not in normalized_text, claim

    assert "1. 자모·한자·약어·혼용어·문맥 삽입" in normalized_text
    assert "1. 탐지 범위:" in normalized_text
    assert "1. 데이터 행 수·스키마·유효성 검사" in normalized_text
    assert "13. 자모·한자" not in normalized_text
    assert "18. 탐지 범위" not in normalized_text
    assert "23. 데이터 행 수" not in normalized_text
    assert "1. 기술 데이터셋의 pii_type" in normalized_text
    assert "4. 기대되는 기술 조치" in normalized_text

    english_abstract_paragraph = document.paragraphs[
        paragraph_texts.index("ABSTRACT") + 2
    ]
    assert english_abstract_paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
    first_reference = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("노희주, 조윤오")
    )
    assert first_reference.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert first_reference.paragraph_format.left_indent
    assert first_reference.paragraph_format.left_indent.mm > 0
    assert first_reference.paragraph_format.first_line_indent
    assert first_reference.paragraph_format.first_line_indent.mm < 0

    table_geometry = []
    for table_index, table in enumerate(document.tables, 1):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        borders = tbl_pr.find(qn("w:tblBorders"))
        assert borders is not None
        for edge in ("top", "bottom", "insideH"):
            assert borders.find(qn(f"w:{edge}")).get(qn("w:val")) == "single"
        for edge in ("left", "right", "insideV"):
            assert borders.find(qn(f"w:{edge}")).get(qn("w:val")) == "nil"
        grid_widths = [
            int(node.get(qn("w:w"))) for node in table._tbl.tblGrid.findall(qn("w:gridCol"))
        ]
        width = int(tbl_w.get(qn("w:w")))
        indent = int(tbl_ind.get(qn("w:w")))
        assert width == sum(grid_widths), (table_index, width, grid_widths)
        assert indent == 120, (table_index, indent)
        assert width == 9970, (table_index, width)
        for row in table.rows:
            for cell_index, cell in enumerate(row.cells):
                tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
                assert int(tc_w.get(qn("w:w"))) == grid_widths[cell_index]
        table_geometry.append(
            {"table": table_index, "columns": len(grid_widths), "width_dxa": width}
        )

    if anonymous:
        with zipfile.ZipFile(path) as archive:
            assert "docProps/custom.xml" not in archive.namelist()
            rsid_count = sum(
                len(re.findall(rb'\bw:rsid\w*="[^"]*"', archive.read(name)))
                for name in archive.namelist()
                if name.endswith(".xml")
            )
            assert rsid_count == 0, rsid_count

    return {
        "file": str(path),
        "anonymous": anonymous,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "sections": len(document.sections),
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "geometry_dxa": actual,
        "table_geometry": table_geometry,
        "structural_qa": "passed",
    }


def main():
    results = [
        verify_document(DOCX, anonymous=False),
        verify_document(ANONYMOUS_DOCX, anonymous=True),
    ]
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
